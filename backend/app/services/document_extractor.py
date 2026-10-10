from io import BytesIO
from pathlib import Path
from docx import Document
from pypdf import PdfReader

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 20000
SUPPORTED = {".pdf", ".docx", ".txt"}

def extract_document(filename: str, data: bytes) -> dict:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED: raise ValueError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")
    if not data: raise ValueError("The uploaded file is empty.")
    if len(data) > MAX_FILE_BYTES: raise ValueError("File is larger than the 5 MB limit.")
    try:
        if ext == ".txt":
            try: text = data.decode("utf-8-sig")
            except UnicodeDecodeError as e: raise ValueError("Could not decode this file as UTF-8 text.") from e
        elif ext == ".pdf":
            reader = PdfReader(BytesIO(data), strict=True)
            if reader.is_encrypted: raise ValueError("Password-protected PDFs are not supported.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            doc = Document(BytesIO(data))
            pieces = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                pieces.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
            text = "\n".join(pieces)
    except ValueError: raise
    except Exception as e: raise ValueError("The file could not be read. It may be corrupted or malformed.") from e
    text = text.strip()
    if not text:
        msg = "No selectable text was found. Scanned PDFs may require OCR, which this app does not provide." if ext == ".pdf" else "No text was found in this document."
        raise ValueError(msg)
    warnings = []
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
        warnings.append("Extracted text was limited to 20,000 characters. Review the preview before processing.")
    return {"filename":Path(filename).name,"file_type":ext[1:].upper(),"extracted_text":text,"word_count":len(text.split()),"character_count":len(text),"warnings":warnings}
