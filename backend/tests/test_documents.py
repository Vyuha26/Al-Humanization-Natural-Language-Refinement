import pytest
from io import BytesIO
from docx import Document
from pypdf import PdfWriter
from app.services.document_extractor import extract_document

def test_txt_extraction():
    result = extract_document("note.txt", "Hello from HumanizeAI".encode())
    assert result["extracted_text"] == "Hello from HumanizeAI"
    assert result["word_count"] == 3

def test_unsupported_file():
    with pytest.raises(ValueError, match="Unsupported"):
        extract_document("payload.exe", b"x")

def test_empty_file():
    with pytest.raises(ValueError, match="empty"):
        extract_document("note.txt", b"")

def test_oversized_file():
    with pytest.raises(ValueError, match="5 MB"):
        extract_document("note.txt", b"a" * (5 * 1024 * 1024 + 1))

def test_docx_paragraph_and_table_extraction():
    doc = Document()
    doc.add_paragraph("A paragraph.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Avery"
    stream = BytesIO()
    doc.save(stream)
    result = extract_document("brief.docx", stream.getvalue())
    assert "A paragraph." in result["extracted_text"]
    assert "Name | Avery" in result["extracted_text"]

def test_corrupted_docx():
    with pytest.raises(ValueError, match="could not be read"):
        extract_document("broken.docx", b"not a Word document")

def test_pdf_with_no_selectable_text_explains_ocr_limit():
    pdf = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(pdf)
    with pytest.raises(ValueError, match="OCR"):
        extract_document("scan.pdf", pdf.getvalue())

def test_text_pdf_extraction():
    stream = b"BT /F1 12 Tf 72 720 Td (Hello from PDF) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    result = extract_document("note.pdf", bytes(pdf))
    assert "Hello from PDF" in result["extracted_text"]

def test_corrupted_pdf():
    with pytest.raises(ValueError, match="could not be read"):
        extract_document("broken.pdf", b"%PDF-not-valid")
